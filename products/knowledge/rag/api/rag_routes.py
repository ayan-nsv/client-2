import os
import uuid
import re
import csv
import tempfile
import zipfile
import shutil
import subprocess
import xml.etree.ElementTree as ET
from typing import List, Optional
from fastapi import APIRouter, UploadFile, File, HTTPException, Depends, status
from fastapi.responses import JSONResponse
from fastapi.encoders import jsonable_encoder
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from sqlalchemy import or_, text
from fastapi import Query
from products.knowledge.rag.utils.pagenation import paginate_query
from fastapi import Request
from shared.database.postgres.database_config import get_db
from products.knowledge.rag.tables.rag_tables import UploadMethod, ExtractedData, CompanyFAQ
from core.tables.user_tables import User
from products.knowledge.rag.schema.rag_schema import (
    UploadFilesResponse,
    UploadPdfPayload,
    UploadUsedFor,
    UploadImagePayload,
    UploadUrlsPayload,
    UploadUrlsResponse,
    UrlsResponse,
    UrlItem,
    DocumentsResponse,
    DocumentsData,
    DocumentItem,
    FAQUploadPayload,
    FAQUploadResponse,
    FAQDetailResponse,
    FAQUpdatePayload,
    FAQListResponse,
    FAQDeletePayload,
    UploadMethodSelectionPayload,
    UploadMethodSelectionResponse,
)
from products.knowledge.rag.services.s3_service import upload_pdf_to_s3, delete_file_from_s3
from shared.logger.log import setup_logger
from shared.utils.error import error, error_handler
from langchain_openai import OpenAIEmbeddings
from langchain_community.document_loaders import PyPDFLoader
from products.knowledge.rag.utils.rag_utils import (
    file_content_hash,
    _image_content_type,
    _extract_text_from_image,
    normalize_text,
    merge_small_chunks,
    chunk_segments_for_rag,
    embed_documents_batched,
    _detect_language_and_translate,
    openai_langchain_kwargs,
    upload_method_used_for_values
)
from products.knowledge.rag.services.url_scraper import scrape_url_for_rag, validate_scrape_url
import pytesseract
from PIL import Image
from shared.utils.auth import auth 
import io

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import xlrd
except ImportError:
    xlrd = None

from core.repository.company_repo import CompanyRepository
from products.market_planner.services.theme_service import ThemeService
from core.repository.user_repo import UserRepository
from products.knowledge.rag.repository.rag_repo import RagRepository
from products.knowledge.rag.utils.rag_query_cache import invalidate_company_cache
from shared.utils.error import error

logger = setup_logger("marketing-app")
router = APIRouter()

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
OPENAI_EMBEDDING_MODEL = os.getenv("OPENAI_EMBEDDING_MODEL") or "text-embedding-3-small"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}

SUPPORTED_DOC_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    # ".odt",
    # ".rtf",
    ".text",
    ".txt",
    ".md",
    ".xlsx",
    ".xls",
    ".och",
}
DOC_CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    # ".odt": "application/vnd.oasis.opendocument.text",
    # ".rtf": "application/rtf",
    ".text": "text/plain",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".och": "application/octet-stream",
}


def _is_static_admin_user(user: dict) -> bool:
    return bool(user.get("is_static_admin"))

def _is_bypass_admin_user(user: dict) -> bool:
    """Static API key, admin secret, or guest special-key auth — no DB user row."""
    return bool(
        user.get("is_static_admin")
        or user.get("is_admin_secret")
        or user.get("is_guest_admin")
    )


def _resolve_payload_user_id(payload_user_id: Optional[str], db: Session) -> Optional[int]:
    """
    Resolve optional request ``user_id`` to internal ``users.id`` PK.

    Matches the rest of the project (user/roles/usage routes):
    - ``users.uuid`` (primary external id), or
    - ``users.firebase_uid`` (Firebase auth uid already stored in DB).

    Integer PK is not accepted as an API ``user_id``.
    """
    if not payload_user_id or not str(payload_user_id).strip():
        return None
    raw = str(payload_user_id).strip()

    # 1) users.uuid
    try:
        user_uuid = uuid.UUID(raw)
        user_record = db.query(User).filter(User.uuid == user_uuid).first()
        if user_record:
            return user_record.id
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"User with uuid '{raw}' not found in database.",
        )
    except ValueError:
        pass

    # 2) users.firebase_uid
    user_record = db.query(User).filter(User.firebase_uid == raw).first()
    if user_record:
        return user_record.id

    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail=(
            f"User '{raw}' not found. Pass users.uuid or users.firebase_uid "
            "(not the integer users.id)."
        ),
    )


def _user_id_for_response(user_int_id: Optional[int], db: Session) -> Optional[str]:
    """Return ``users.uuid`` string — same external user id used across the project."""
    if user_int_id is None:
        return None
    user_record = db.query(User).filter(User.id == user_int_id).first()
    if not user_record:
        return None
    return str(user_record.uuid)


def _reset_upload_methods_sequence(db: Session) -> None:
    """Reset upload_methods PK sequence when it drifts behind max(id)."""
    db.execute(
        text(
            "SELECT setval("
            "  pg_get_serial_sequence('upload_methods', 'id'),"
            "  COALESCE((SELECT MAX(id) FROM upload_methods), 0) + 1,"
            "  false"
            ")"
        )
    )


def _find_url_upload_methods(
    db: Session,
    company_id: int,
    used_for: str,
    *candidate_urls: str,
) -> List[UploadMethod]:
    """Existing URL sources for this company/channel matching any candidate URL."""
    urls = [u for u in candidate_urls if u]
    if not urls:
        return []
    return (
        db.query(UploadMethod)
        .filter(
            UploadMethod.company_id == company_id,
            UploadMethod.method_type == "url",
            UploadMethod.used_for == used_for,
            or_(
                UploadMethod.url.in_(urls),
                UploadMethod.file_path.in_(urls),
            ),
        )
        .order_by(UploadMethod.created_at.desc())
        .all()
    )


def _existing_content_hash_for_upload(
    db: Session, upload_method_id: int
) -> Optional[str]:
    row = (
        db.query(ExtractedData.content_hash)
        .filter(ExtractedData.upload_method_id == upload_method_id)
        .first()
    )
    return row[0] if row else None


def _find_company_upload_by_content_hash(
    db: Session,
    company_id: int,
    content_hash: str,
    exclude_upload_method_id: Optional[int] = None,
) -> Optional[UploadMethod]:
    """Another knowledge source for this company that already indexes the same content."""
    query = (
        db.query(UploadMethod)
        .join(ExtractedData, ExtractedData.upload_method_id == UploadMethod.id)
        .filter(
            ExtractedData.company_id == company_id,
            ExtractedData.content_hash == content_hash,
        )
    )
    if exclude_upload_method_id is not None:
        query = query.filter(UploadMethod.id != exclude_upload_method_id)
    return query.order_by(UploadMethod.created_at.desc()).first()


if not OPENAI_API_KEY:
    logger.warning("OPENAI_API_KEY not set. PDF embedding extraction will fail.")


def _document_content_type(filename: str) -> str:
    ext = os.path.splitext((filename or "").lower())[1]
    return DOC_CONTENT_TYPES.get(ext, "application/octet-stream")


def _decode_text_bytes(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except Exception:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _tabular_text_from_csv_bytes(file_bytes: bytes) -> str:
    """
    Parse bytes as CSV/TSV-style text (comma, tab, or semicolon separated).
    Used when a file has a .xls extension but is actually plain-text (common export mistake).
    """
    raw = _decode_text_bytes(file_bytes)
    if not raw.strip():
        return ""
    sample = raw[:8192]
    delimiters = [",", "\t", ";"]
    best_rows: Optional[List[List[str]]] = None
    best_d = ","
    for d in delimiters:
        try:
            rows = list(csv.reader(io.StringIO(raw), delimiter=d))
            nonempty = sum(1 for r in rows if any((c or "").strip() for c in r))
            if nonempty == 0:
                continue
            if best_rows is None or nonempty > sum(
                1 for r in best_rows if any((c or "").strip() for c in r)
            ):
                best_rows = rows
                best_d = d
        except Exception:
            continue
    if not best_rows:
        return raw.strip()
    lines: List[str] = []
    for row in best_rows:
        if any((c or "").strip() for c in row):
            lines.append("\t".join(row))
    return "\n".join(lines)


def _open_xlsx_workbook(file_bytes: bytes):
    if load_workbook is None:
        raise ValueError(
            "openpyxl is required for .xlsx files. Install with: pip install openpyxl"
        )
    return load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)


def _extract_text_xlsx(file_bytes: bytes) -> str:
    wb = _open_xlsx_workbook(file_bytes)
    lines: List[str] = []
    try:
        for sheet in wb.worksheets:
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(x.strip() for x in cells):
                    lines.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)


def _open_xls_workbook(file_bytes: bytes):
    if xlrd is None:
        raise ValueError(
            "xlrd is required for binary .xls files. Install with: pip install xlrd"
        )

    try:
        book = xlrd.open_workbook(file_contents=file_bytes)
    except Exception as e:
        return None, e
    return book, None


def _extract_text_xls(file_bytes: bytes) -> str:
    """
    True Excel 97-2003 (.xls BIFF) via xlrd. Many files named .xls are actually CSV;
    if xlrd fails (e.g. Expected BOF record), fall back to CSV/TSV parsing.
    """
    book, error = _open_xls_workbook(file_bytes)
    if book is None:
        fallback = _tabular_text_from_csv_bytes(file_bytes)
        if fallback.strip() and len(fallback.strip()) > 10:
            logger.info(
                "xlrd could not open .xls as binary workbook; using CSV/TSV parse instead: %s",
                error,
            )
            return fallback
        raise ValueError(
            f"Could not read as Excel .xls or as CSV/TSV text: {error}. "
            "Try saving as real Excel 97-2003 (.xls), .xlsx, or plain .csv."
        ) from error

    lines: List[str] = []
    for sheet in book.sheets():
        lines.append(f"## {sheet.name}")
        for row_idx in range(sheet.nrows):
            row = sheet.row(row_idx)
            cells = ["" if c.ctype == xlrd.XL_CELL_EMPTY else str(c.value) for c in row]
            if any(str(x).strip() for x in cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


def _extract_text_docx(file_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        xml_data = zf.read("word/document.xml")
    root = ET.fromstring(xml_data)
    parts = [t.text for t in root.iter() if t.tag.endswith("}t") and t.text]
    return "\n".join(parts)


# def _extract_text_odt(file_bytes: bytes) -> str:
#     with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
#         xml_data = zf.read("content.xml")
#     root = ET.fromstring(xml_data)
#     return "".join(root.itertext())


def _extract_text_doc(file_bytes: bytes) -> str:
    antiword = shutil.which("antiword")
    if antiword:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".doc") as tmp:
            tmp.write(file_bytes)
            path = tmp.name
        try:
            proc = subprocess.run(
                [antiword, path],
                capture_output=True,
                text=True,
                timeout=30,
                check=False,
            )
            if proc.returncode == 0 and proc.stdout.strip():
                return proc.stdout
        finally:
            try:
                os.unlink(path)
            except Exception:
                pass
    raise ValueError(
        "Could not extract text from .doc file. Install system package "
        "'antiword' on the server, or upload .docx instead."
    )


# def _strip_rtf(raw: str) -> str:
#     text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
#     text = re.sub(r"\\[a-zA-Z]+\s*-?\d*\s?", " ", text)
#     text = text.replace("{", " ").replace("}", " ")
#     return re.sub(r"\s+", " ", text).strip()


def _extract_document_chunks(filename: str, file_bytes: bytes) -> List[str]:
    ext = os.path.splitext((filename or "").lower())[1]

    if ext == ".pdf":
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(file_bytes)
            path = tmp.name
        try:
            loader = PyPDFLoader(path)
            docs = loader.load()
            return [p.page_content for p in docs if (p.page_content or "").strip()]
        finally:
            try:
                os.unlink(path)
            except Exception:
                pass

    elif ext == ".doc":
        text = _extract_text_doc(file_bytes)
    elif ext == ".docx":
        text = _extract_text_docx(file_bytes)
    elif ext in {".text", ".txt", ".md", ".och"}:
        text = _decode_text_bytes(file_bytes)
    elif ext == ".xlsx":
        text = _extract_text_xlsx(file_bytes)
    elif ext == ".xls":
        text = _extract_text_xls(file_bytes)
    else:
        raise ValueError(f"Unsupported extension: {ext}")

    text = (text or "").strip()
    return [text] if text else []

@router.post(
    "/upload-faq",
    response_model=FAQUploadResponse,
    response_model_exclude_none=True,
    tags=["rag"],
)
async def upload_faq(
    payload: FAQUploadPayload,
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user)
    #api_key: str = Depends(api_key_security)

):
    """
    Save FAQ questions and answers for a company. Questions are embedded for similarity search.
    """
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")

        firebase_uid = user["uid"]

        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(payload.company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        if not _is_static_admin_user(user) and not UserRepository(db).is_user_admin(firebase_uid, company_int_id):
            return error.PermissionDenied(
                message="Admin privileges required"
            )

        if not OPENAI_API_KEY:
            return error.InternalServerError(message="OPENAI_API_KEY not configured.")

        embedder = OpenAIEmbeddings(
            api_key=OPENAI_API_KEY,
            model=OPENAI_EMBEDDING_MODEL,
            **openai_langchain_kwargs(),
        )

        if not payload.faqs:
            return error.BadRequest(message="At least one FAQ (question + answer) is required")

        # Enforce category to be 'faq' unless explicitly set to 'common_faq'
        actual_category = "common_faq" if payload.category == "common_faq" else "faq"

        if actual_category == "common_faq":
            # Count only one language to get the true number of FAQ items
            existing_count = db.query(CompanyFAQ).filter(
                CompanyFAQ.company_id == company_int_id,
                CompanyFAQ.category == "common_faq",
                CompanyFAQ.language == "sv"  # Or 'en', just need one
            ).count()
            
            if existing_count + len(payload.faqs) > 10:
                return error.BadRequest(message=f"Maximum 10 common_faq allowed per company. You already have {existing_count} and are trying to add {len(payload.faqs)}.")

        faq_records: List[dict] = []

        # Detect language, translate, and prepare FAQ records
        for item in payload.faqs:
            q_id = uuid.uuid4().hex
            original_question = (item.question or "").strip()
            original_answer = (item.answer or "").strip()

            if not original_question or not original_answer:
                return error.BadRequest(message="Each FAQ must include a non-empty question and answer.")

            translation_result = _detect_language_and_translate(
                question=original_question,
                answer=original_answer,
            )

            source_language = translation_result["source_language"]
            target_language = translation_result["target_language"]
            translated_question = translation_result["translated_question"]
            translated_answer = translation_result["translated_answer"]

            # If language is not English or Swedish, abort and roll back everything
            if source_language not in ("en", "sv"):
                return error.BadRequest(message="Unknown language, Allowed languages are (English / Swedish)")

            if not translated_question or not translated_answer:
                return error.InternalServerError(message="Translation failed for the provided FAQ.")

            if target_language not in ("en", "sv"):
                return error.InternalServerError(message="Translation did not return a valid target language.")

            # Store the original question/answer and detected language
            faq_records.append(
                {
                    "question": original_question,
                    "answer": original_answer,
                    "language": source_language,
                    "question_id": q_id,
                }
            )

            # Also store the translated version as a separate FAQ row
            faq_records.append(
                {
                    "question": translated_question,
                    "answer": translated_answer,
                    "language": target_language,
                    "question_id": q_id,
                }
            )

        # Create embeddings for all stored FAQs (original + translated)
        texts_for_embedding: List[str] = [
            f"{record['question']}\n\n{record['answer']}" for record in faq_records
        ]
        embeddings = embedder.embed_documents(texts_for_embedding)

        for record, vector in zip(faq_records, embeddings):
            faq = CompanyFAQ(
                company_id=company_int_id,
                question=record["question"],
                answer=record["answer"],
                embedding=vector,
                language=record["language"],
                category=actual_category,
                question_id=record["question_id"]
            )
            db.add(faq)

        db.commit()
        return {
            "success": True,
            "message": "FAQ saved successfully",
            "data": {"saved_count": len(faq_records)},
        }
    except HTTPException as exc:
        db.rollback()
        return error.InternalServerError(message=str(exc.detail))
    except Exception as e:
        db.rollback()
        logger.error(f"FAQ upload failed: {str(e)}")
        return error.InternalServerError(message=f"FAQ upload failed: {str(e)}")


@router.get(
    "/{company_id}/faqs",
    response_model=FAQListResponse,
    response_model_exclude_none=True,
    tags=["rag"]
)
async def list_faqs(
    request: Request,
    company_id: str,
    category: Optional[str] = Query(None, description="Filter by category (e.g. common_faq)"),
    language: Optional[str] = Query(None, description="Filter by language (e.g. en, sv)"),
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    #api_key: str = Depends(api_key_security)

):
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        base_query = db.query(CompanyFAQ).filter(CompanyFAQ.company_id == company_int_id)
        
        if category:
            base_query = base_query.filter(CompanyFAQ.category == category)
        if language:
            base_query = base_query.filter(CompanyFAQ.language == language)
            
        base_query = base_query.order_by(CompanyFAQ.created_at.desc())

        faqs, pagination = paginate_query(
            base_query,
            request,
            page,
            page_size
        )

        faq_list = [
            {
                "uuid": faq.uuid,
                "question": faq.question,
                "answer": faq.answer,
                "category": faq.category,
                "language": faq.language
            }
            for faq in faqs
        ]

        return {
            "success": True,
            "message": "FAQs retrieved successfully",
            "data": {
                "faqs": faq_list,
                "pagination": pagination
            }
        }
    except ValueError as ve:
        return {"success": False, "error": str(ve)}
    except Exception as e:
        logger.error(f"Failed to retrieve FAQs: {str(e)}")
        return {"success": False, "error": "Failed to retrieve FAQs"}


@router.get(
    "/faqs/{faq_uuid}",
    response_model=FAQDetailResponse,
    response_model_exclude_none=True,
    tags=["rag"]
)
async def get_faq(
    faq_uuid: str,
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    #api_key: str = Depends(api_key_security)

):
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        faq = db.query(CompanyFAQ).filter(CompanyFAQ.uuid == faq_uuid).first()
        if not faq:
            return error.NotFound(message="FAQ not found")
            
        return {
            "success": True,
            "data": {
                "uuid": faq.uuid,
                "question": faq.question,
                "answer": faq.answer,
                "category": faq.category,
                "language": faq.language
            }
        }
    except Exception as e:
        logger.error(f"Failed to fetch FAQ: {str(e)}")
        return error.InternalServerError(message=f"Failed to fetch FAQ: {str(e)}")


@router.put(
    "/faqs/{faq_uuid}",
    response_model=FAQDetailResponse,
    response_model_exclude_none=True,
    tags=["rag"]
)
async def update_faq(
    faq_uuid: str,
    payload: FAQUpdatePayload,
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    #api_key: str = Depends(api_key_security)

):
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        firebase_uid = user["uid"]

        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(payload.company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        if not _is_static_admin_user(user) and not UserRepository(db).is_user_admin(firebase_uid, company_int_id):
            return error.Forbidden(message="Admin privileges required")

        faq = db.query(CompanyFAQ).filter(
            CompanyFAQ.uuid == faq_uuid,
            CompanyFAQ.company_id == company_int_id
        ).first()
        
        if not faq:
            return error.NotFound(message="FAQ not found")

        # Update fields if provided
        if payload.question is not None:
            faq.question = payload.question
        if payload.answer is not None:
            faq.answer = payload.answer
        if payload.category is not None:
            faq.category = payload.category
            
        # Re-embed if question or answer changed
        if payload.question is not None or payload.answer is not None:
            if not OPENAI_API_KEY:
                return error.InternalServerError(message="OPENAI_API_KEY not configured.")

            embedder = OpenAIEmbeddings(
                api_key=OPENAI_API_KEY,
                model=OPENAI_EMBEDDING_MODEL,
                **openai_langchain_kwargs(),
            )

            # Re-translate the updated FAQ
            translation_result = _detect_language_and_translate(
                question=faq.question,
                answer=faq.answer,
            )
            
            # Embed the updated FAQ
            text_for_embedding = f"{faq.question}\n\n{faq.answer}"
            embeddings = embedder.embed_documents([text_for_embedding])
            
            faq.embedding = embeddings[0]
            faq.language = translation_result["source_language"]
            
            # Sync the partner FAQ if it exists
            if faq.question_id:
                partner_faq = db.query(CompanyFAQ).filter(
                    CompanyFAQ.question_id == faq.question_id,
                    CompanyFAQ.uuid != faq.uuid,
                    CompanyFAQ.company_id == company_int_id
                ).first()
                
                if partner_faq:
                    partner_faq.question = translation_result["translated_question"]
                    partner_faq.answer = translation_result["translated_answer"]
                    partner_faq.language = translation_result["target_language"]
                    
                    partner_text = f"{partner_faq.question}\n\n{partner_faq.answer}"
                    partner_embeddings = embedder.embed_documents([partner_text])
                    partner_faq.embedding = partner_embeddings[0]

        db.commit()
        db.refresh(faq)
        
        return {
            "success": True,
            "data": {
                "uuid": faq.uuid,
                "question": faq.question,
                "answer": faq.answer,
                "category": faq.category,
                "language": faq.language
            }
        }
    except Exception as e:
        db.rollback()
        logger.error(f"Failed to update FAQ: {str(e)}")
        return error.InternalServerError(message=f"Failed to update FAQ: {str(e)}")



@router.delete(
    "/faqs/{faq_uuid}",
    response_model=FAQDetailResponse,
    response_model_exclude_none=True,
    tags=["rag"]
)
async def delete_faq(
    faq_uuid: str,
    payload: FAQDeletePayload,
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    #api_key: str = Depends(api_key_security)

):
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        firebase_uid = user["uid"]

        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(payload.company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        if not _is_static_admin_user(user) and not UserRepository(db).is_user_admin(firebase_uid, company_int_id):
            return error.Forbidden(message="Admin privileges required")

        # Find the specific FAQ to delete
        faq = db.query(CompanyFAQ).filter(
            CompanyFAQ.uuid == faq_uuid,
            CompanyFAQ.company_id == company_int_id
        ).first()
        
        if not faq:
            return error.NotFound(message="FAQ not found")

        # Use question_id to delete all linked versions (EN and SV)
        if faq.question_id:
            deleted_count = db.query(CompanyFAQ).filter(
                CompanyFAQ.question_id == faq.question_id,
                CompanyFAQ.company_id == company_int_id
            ).delete()
            logger.info(f"Deleted {deleted_count} FAQs for question_id: {faq.question_id}")
        else:
            # Fallback if question_id is somehow missing
            db.delete(faq)
            logger.info(f"Deleted single FAQ: {faq_uuid}")

        db.commit()
        
        # Invalidate company RAG cache so changes are reflected in the agent immediately
        await invalidate_company_cache(payload.company_id)
        
        return {
            "success": True,
            "message": "FAQ and its translations deleted successfully"
        }
    except Exception as e:
        db.rollback()
        logger.error(f"Failed to delete FAQ: {str(e)}")
        return error.InternalServerError(message=f"Failed to delete FAQ: {str(e)}")


@router.post("/upload-files", 
             response_model=UploadFilesResponse, 
             response_model_exclude_none=True,
             tags=["rag"])
async def upload_files(
    files: List[UploadFile] = File(...),
    payload: UploadPdfPayload = Depends(UploadPdfPayload.as_form),
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    #api_key: str = Depends(api_key_security)

):
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        if len(files) > 5:
            return error.BadRequest(message="Maximum 5 files allowed per request")

        if len(files) == 0:
            return error.BadRequest(message="No files uploaded")

        for file in files:
            ext = os.path.splitext((file.filename or "").lower())[1]
            if ext not in SUPPORTED_DOC_EXTENSIONS:
                return error.BadRequest(message=f"File {file.filename} is not allowed (use .pdf, .doc, .docx, .text, .txt, .md, .xlsx, .xls, .och)")

        used_for = (
            payload.used_for.value
            if payload.used_for is not None
            else UploadUsedFor.voice_call.value
        )
        firebase_uid = user["uid"]
        user_int_id = None
        if not _is_static_admin_user(user):
            user_record = db.query(User).filter(User.firebase_uid == firebase_uid).first()
            if not user_record:
                return error.NotFound(message=f"User with Firebase UID '{firebase_uid}' not found in database.")
            # Convert UUIDs to integer IDs (custom response format on errors)
            try:
                user_int_id = UserRepository(db).get_user_id_from_firebase_uid(firebase_uid)
            except HTTPException as exc:
                return error.InternalServerError(message=str(exc.detail))

        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(payload.company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        if not _is_static_admin_user(user) and not UserRepository(db).is_user_admin(firebase_uid, company_int_id):
            return error.Forbidden(message="Admin privileges required")

        
        if not OPENAI_API_KEY:
            return error.InternalServerError(message="OPENAI_API_KEY not configured.")

        embedder = OpenAIEmbeddings(
            api_key=OPENAI_API_KEY,
            model=OPENAI_EMBEDDING_MODEL,
            **openai_langchain_kwargs(),
        )

        uploaded_files = []
        errors = []

        for file in files:
            try:
                file_bytes = await file.read()
                filename = file.filename or "unknown.pdf"

                file_uuid = uuid.uuid4().hex
                s3_file_path = f"documents/{company_int_id}/{file_uuid}_{filename}"

                full_url = await upload_pdf_to_s3(
                    file_bytes=file_bytes,
                    file_path=s3_file_path,
                    content_type=_document_content_type(filename),
                )

                upload_method = UploadMethod(
                    uploaded_by_id=user_int_id,
                    company_id=company_int_id,
                    method_type="attachment",
                    filename=filename,
                    file_path=full_url,
                    url=full_url,
                    used_for=used_for,
                )

                db.add(upload_method)
                db.flush()

                raw_chunks = _extract_document_chunks(filename, file_bytes)
                content_hash = file_content_hash(file_bytes)

                # Remove old chunks for this upload
                db.query(ExtractedData).filter(
                    ExtractedData.upload_method_id == upload_method.id
                ).delete()

                # STEP 1: Normalize text
                cleaned_chunks = [normalize_text(c) for c in raw_chunks]

                # Remove empty chunks
                cleaned_chunks = [c for c in cleaned_chunks if c]

                # STEP 2: Merge small pages/fragments (word-based)
                cleaned_chunks = merge_small_chunks(
                    cleaned_chunks,
                    min_words=5
                )

                # STEP 3: Bounded overlapping chunks for faster, sharper retrieval
                cleaned_chunks = chunk_segments_for_rag(cleaned_chunks)

                if not cleaned_chunks:
                    raise ValueError("No extractable text from this document.")

                # STEP 4: Embed in API-sized batches
                embeddings = embed_documents_batched(embedder, cleaned_chunks)

                for i, (text, vector) in enumerate(zip(cleaned_chunks, embeddings)):
                    extracted_data = ExtractedData(
                        upload_method_id=upload_method.id,
                        chunk_index=i,
                        chunk_content=text,
                        company_id=company_int_id,
                        content_hash=content_hash,
                        embedding=vector
                    )
                    db.add(extracted_data)

                uploaded_files.append({"file_path": full_url})
                logger.info(f"Successfully processed document: {filename}")

            except Exception as e:
                error_msg = f"Failed to process file {file.filename}: {str(e)}"
                logger.error(error_msg)
                errors.append(error_msg)

        if uploaded_files:
            db.commit()
            await invalidate_company_cache(payload.company_id)
        else:
            db.rollback()
            raise HTTPException(
                status_code=500,
                detail=f"Failed to upload files. Errors: {errors}"
            )

        if errors:
            return JSONResponse(
                status_code=207,
                content={
                    "success": True,
                    "data": {"uploaded_files": uploaded_files},
                    "error": f"Some files failed: {', '.join(errors)}"
                }
            )

        return {
            "success": True,
            "message": "Files uploaded successfully",
            "data": {"uploaded_files": uploaded_files}
        }

    except HTTPException as exc:
        db.rollback()
        return error.InternalServerError(message=str(exc.detail))

    except Exception as e:
        db.rollback()
        logger.error(f"Unexpected error: {str(e)}")
        return error.InternalServerError(message=f"Files upload failed: {str(e)}")


@router.post(
    "/upload-urls",
    response_model=UploadUrlsResponse,
    response_model_exclude_none=True,
    tags=["rag"],
)
async def upload_urls(
    payload: UploadUrlsPayload,
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
):
    """
    Ingest important URLs into the company RAG index.

    Re-submitting the same URL for the same ``used_for`` channel re-processes it
    (upsert). Unchanged content (same ``content_hash``) is skipped. Duplicate
    page content already indexed for the company is also skipped.
    """
    try:
        raw_urls = [u.strip() for u in payload.urls if (u or "").strip()]

        if len(raw_urls) > 5:
            return error.BadRequest(message="Maximum 5 URLs allowed per request")

        if len(raw_urls) == 0:
            return error.BadRequest(message="No URLs provided")

        # Deduplicate within the request (same pattern as other multi-item APIs)
        urls: List[str] = []
        seen_request_urls = set()
        for url in raw_urls:
            try:
                normalized = validate_scrape_url(url)
            except ValueError as exc:
                return error.BadRequest(message=f"Invalid URL {url}: {exc}")
            key = normalized.rstrip("/").lower()
            if key in seen_request_urls:
                continue
            seen_request_urls.add(key)
            urls.append(normalized)

        used_for = (
            payload.used_for.value
            if payload.used_for is not None
            else UploadUsedFor.voice_call.value
        )
        firebase_uid = user["uid"]
        user_int_id = None

        # Same pattern as upload-files, plus bypass for admin-secret / guest auth
        if _is_bypass_admin_user(user):
            try:
                user_int_id = _resolve_payload_user_id(payload.user_id, db)
            except HTTPException as exc:
                return error.InternalServerError(message=str(exc.detail))
        else:
            user_record = db.query(User).filter(User.firebase_uid == firebase_uid).first()
            if not user_record:
                return error.InternalServerError(message=f"User with Firebase UID '{firebase_uid}' not found in database.")
            try:
                user_int_id = ThemeService(db).get_user_id_for_usage_tracking(firebase_uid)
            except HTTPException as exc:
                return error.InternalServerError(message=str(exc.detail))
            if payload.user_id and str(payload.user_id).strip():
                try:
                    user_int_id = _resolve_payload_user_id(payload.user_id, db)
                except HTTPException as exc:
                    return error.InternalServerError(message=str(exc.detail))

        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(payload.company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        if not _is_bypass_admin_user(user) and not UserRepository(db).is_user_admin(firebase_uid, company_int_id):
            return error.Forbidden(message="Admin privileges required")

        if not OPENAI_API_KEY:
            return error.InternalServerError(message="OPENAI_API_KEY not configured.")

        embedder = OpenAIEmbeddings(
            api_key=OPENAI_API_KEY,
            model=OPENAI_EMBEDDING_MODEL,
            **openai_langchain_kwargs(),
        )

        results = []
        errors = []
        any_index_change = False
        response_user_id = _user_id_for_response(user_int_id, db)

        for url in urls:
            try:
                scraped = await scrape_url_for_rag(url)
                normalized_url = scraped["url"]
                filename = scraped["filename"]
                raw_chunks = scraped["segments"]
                content_hash = file_content_hash(
                    "\n\n".join(raw_chunks).encode("utf-8")
                )

                existing_rows = _find_url_upload_methods(
                    db, company_int_id, used_for, url, normalized_url
                )
                upload_method = existing_rows[0] if existing_rows else None

                # Collapse accidental duplicate URL rows from older uploads
                for extra in existing_rows[1:]:
                    db.delete(extra)
                    any_index_change = True

                # Same page content already indexed for this company (another source)
                if upload_method is None:
                    dup_source = _find_company_upload_by_content_hash(
                        db, company_int_id, content_hash
                    )
                    if dup_source is not None:
                        results.append({
                            "uuid": dup_source.uuid,
                            "url": normalized_url,
                            "user_id": response_user_id,
                            "status": "skipped_duplicate_content",
                        })
                        logger.info(
                            "Skipped duplicate content for %s (matches upload_method=%s)",
                            normalized_url,
                            dup_source.id,
                        )
                        continue

                # Re-process path: unchanged content → skip re-embed
                if upload_method is not None:
                    existing_hash = _existing_content_hash_for_upload(
                        db, upload_method.id
                    )
                    if existing_hash and existing_hash == content_hash:
                        upload_method.filename = filename
                        upload_method.file_path = normalized_url
                        upload_method.url = normalized_url
                        if user_int_id is not None:
                            upload_method.uploaded_by_id = user_int_id
                        results.append({
                            "uuid": upload_method.uuid,
                            "url": normalized_url,
                            "user_id": response_user_id,
                            "status": "skipped_unchanged",
                        })
                        logger.info(
                            "Skipped unchanged URL content: %s", normalized_url
                        )
                        continue

                cleaned_chunks = [normalize_text(c) for c in raw_chunks]
                cleaned_chunks = [c for c in cleaned_chunks if c]
                cleaned_chunks = merge_small_chunks(cleaned_chunks, min_words=5)
                cleaned_chunks = chunk_segments_for_rag(cleaned_chunks)

                if not cleaned_chunks:
                    raise ValueError("No extractable text from this URL.")

                embeddings = embed_documents_batched(embedder, cleaned_chunks)
                is_update = upload_method is not None

                def _persist_url_upload() -> UploadMethod:
                    nonlocal upload_method
                    if upload_method is None:
                        upload_method = UploadMethod(
                            uploaded_by_id=user_int_id,
                            company_id=company_int_id,
                            method_type="url",
                            filename=filename,
                            file_path=normalized_url,
                            url=normalized_url,
                            used_for=used_for,
                        )
                        db.add(upload_method)
                        db.flush()
                    else:
                        upload_method.filename = filename
                        upload_method.file_path = normalized_url
                        upload_method.url = normalized_url
                        upload_method.used_for = used_for
                        if user_int_id is not None:
                            upload_method.uploaded_by_id = user_int_id
                        db.flush()

                    db.query(ExtractedData).filter(
                        ExtractedData.upload_method_id == upload_method.id
                    ).delete()

                    for i, (text_chunk, vector) in enumerate(
                        zip(cleaned_chunks, embeddings)
                    ):
                        db.add(
                            ExtractedData(
                                upload_method_id=upload_method.id,
                                chunk_index=i,
                                chunk_content=text_chunk,
                                company_id=company_int_id,
                                content_hash=content_hash,
                                embedding=vector,
                            )
                        )
                    return upload_method

                try:
                    with db.begin_nested():
                        upload_method = _persist_url_upload()
                except IntegrityError as integrity_exc:
                    logger.warning(
                        "upload_methods insert conflict for %s; resetting sequence: %s",
                        normalized_url,
                        integrity_exc,
                    )
                    _reset_upload_methods_sequence(db)
                    upload_method = None if not is_update else upload_method
                    with db.begin_nested():
                        upload_method = _persist_url_upload()

                any_index_change = True
                status_label = "updated" if is_update else "created"
                results.append({
                    "uuid": upload_method.uuid,
                    "url": normalized_url,
                    "user_id": response_user_id,
                    "status": status_label,
                })
                logger.info(
                    "Successfully %s URL: %s", status_label, normalized_url
                )

            except HTTPException as exc:
                error_msg = str(exc.detail)
                logger.error("Failed to process URL %s: %s", url, error_msg)
                errors.append(f"Failed to process URL {url}: {error_msg}")
                results.append({
                    "url": url,
                    "status": "failed",
                    "error": error_msg,
                    "user_id": response_user_id,
                })
            except Exception as e:
                error_msg = str(e)
                logger.error("Failed to process URL %s: %s", url, error_msg)
                errors.append(f"Failed to process URL {url}: {error_msg}")
                results.append({
                    "url": url,
                    "status": "failed",
                    "error": error_msg,
                    "user_id": response_user_id,
                })

        succeeded = [r for r in results if r.get("status") != "failed"]
        if succeeded or any_index_change:
            db.commit()
            if any_index_change:
                await invalidate_company_cache(payload.company_id)
        else:
            db.rollback()
            raise HTTPException(
                status_code=500,
                detail=f"Failed to scrape URLs. Errors: {errors}",
            )

        if errors:
            return JSONResponse(
                status_code=207,
                content=jsonable_encoder({
                    "success": True,
                    "message": "Some URLs could not be processed",
                    "data": {"uploaded_urls": results},
                    "error": f"Some URLs failed: {', '.join(errors)}",
                }),
            )

        return {
            "success": True,
            "message": "URLs scraped and indexed successfully",
            "data": {"uploaded_urls": results},
        }

    except HTTPException as exc:
        db.rollback()
        return error.InternalServerError(message=str(exc.detail))

    except Exception as e:
        db.rollback()
        logger.error(f"Unexpected error during URL upload: {str(e)}")
        return error.InternalServerError(message=f"URL upload failed: {str(e)}")


@router.post("/upload-images",
             response_model=UploadFilesResponse,
             response_model_exclude_none=True,
             tags=["rag"])
async def upload_images(
    files: List[UploadFile] = File(...),
    payload: UploadImagePayload = Depends(UploadImagePayload.as_form),
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    # api_key: str = Depends(api_key_security)

):
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        if len(files) > 5:
            return error.BadRequest(message="Maximum 5 image files allowed per request")

        if len(files) == 0:
            return error.BadRequest(message="No files uploaded")

        for file in files:
            ext = os.path.splitext((file.filename or "").lower())[1]
            if ext not in IMAGE_EXTENSIONS:
                return error.BadRequest(message=f"File {file.filename} is not an allowed image (use .jpg, .jpeg, .png, .webp)")

        firebase_uid = user["uid"]
        user_int_id = None
        if not _is_static_admin_user(user):
            user_record = db.query(User).filter(User.firebase_uid == firebase_uid).first()
            if not user_record:
                return error.NotFound(message=f"User with Firebase UID '{firebase_uid}' not found in database.")
                   
            try:
                user_int_id = UserRepository(db).get_user_id_from_firebase_uid(firebase_uid)
            except HTTPException as exc:
                return error.InternalServerError(message=str(exc.detail))

        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(payload.company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        if not _is_static_admin_user(user) and not UserRepository(db).is_user_admin(firebase_uid, company_int_id):
            return error.Forbidden(message="Admin privileges required")

        if not OPENAI_API_KEY:
            return error.InternalServerError(message="OPENAI_API_KEY not configured.")

        embedder = OpenAIEmbeddings(
            api_key=OPENAI_API_KEY,
            model=OPENAI_EMBEDDING_MODEL,
            **openai_langchain_kwargs(),
        )

        uploaded_files = []
        errors = []

        for file in files:
            try:
                file_bytes = await file.read()
                filename = file.filename or "unknown.jpg"

                file_uuid = uuid.uuid4().hex
                s3_file_path = f"images/{company_int_id}/{file_uuid}_{filename}"
                content_type = _image_content_type(filename)

                full_url = await upload_pdf_to_s3(
                    file_bytes=file_bytes,
                    file_path=s3_file_path,
                    content_type=content_type
                )

                upload_method = UploadMethod(
                    uploaded_by_id=user_int_id,
                    company_id=company_int_id,
                    method_type="attachment",
                    filename=filename,
                    file_path=full_url,
                    url=full_url
                )

                db.add(upload_method)
                db.flush()

                # Extract text from image (OCR)
                raw_chunks = _extract_text_from_image(file_bytes)

                content_hash = file_content_hash(file_bytes)

                db.query(ExtractedData).filter(
                    ExtractedData.upload_method_id == upload_method.id
                ).delete()

                # Normalize text
                cleaned_chunks = [normalize_text(c) for c in raw_chunks]

                # Remove empty chunks
                cleaned_chunks = [c for c in cleaned_chunks if c]

                # Merge small OCR fragments
                cleaned_chunks = merge_small_chunks(cleaned_chunks, min_words=5)

                cleaned_chunks = chunk_segments_for_rag(cleaned_chunks)

                if cleaned_chunks:
                    embeddings = embed_documents_batched(embedder, cleaned_chunks)
                else:
                    embeddings = []

                for i, (text, vector) in enumerate(zip(cleaned_chunks, embeddings)):
                    extracted_data = ExtractedData(
                        upload_method_id=upload_method.id,
                        chunk_index=i,
                        chunk_content=text,
                        company_id=company_int_id,
                        content_hash=content_hash,
                        embedding=vector
                    )
                    db.add(extracted_data)

                uploaded_files.append({"file_path": full_url})
                logger.info(f"Successfully processed image: {filename}")

            except Exception as e:
                error_msg = f"Failed to process file {file.filename}: {str(e)}"
                logger.error(error_msg)
                errors.append(error_msg)

        if uploaded_files:
            db.commit()
        else:
            db.rollback()
            return error.InternalServerError(message=f"Failed to upload files. Errors: {errors}")
             
        if errors:
            return JSONResponse(
                status_code=207,
                content={
                    "success": True,
                    "data": {"uploaded_files": uploaded_files},
                    "error": f"Some files failed: {', '.join(errors)}"
                }
            )

        return {
            "success": True,
            "message": "Files uploaded successfully",
            "data": {"uploaded_files": uploaded_files}
        }

    except HTTPException as exc:
        db.rollback()
        return error.InternalServerError(message=str(exc.detail))

    except Exception as e:
        db.rollback()
        logger.error(f"Unexpected error: {str(e)}")
        return error.InternalServerError(message=f"Image upload failed: {str(e)}")
           


@router.patch(
    "/upload-methods/{upload_method_uuid}",
    response_model=UploadMethodSelectionResponse,
    response_model_exclude_none=True,
    tags=["rag"],
)
async def update_upload_method_selection(
    upload_method_uuid: str,
    payload: UploadMethodSelectionPayload,
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    # api_key: str = Depends(api_key_security)

):
    """Set ``upload_methods.is_selected`` for a knowledge-base document."""
    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        try:
            doc_uuid = uuid.UUID(upload_method_uuid)
        except ValueError:
            return error.BadRequest(message="Invalid upload method UUID")
              

        upload_method = RagRepository(db).get_upload_method_by_uuid(doc_uuid)
        upload_method.is_selected = payload.is_selected
        db.commit()
        db.refresh(upload_method)

        if upload_method.company_id is not None:
            company_uuid = CompanyRepository(db).get_company_uuid_from_id(upload_method.company_id)
            await invalidate_company_cache(company_uuid)

        return {
            "success": True,
            "message": "Upload method selection updated",
            "data": {
                "id": upload_method.id,
                "uuid": upload_method.uuid,
                "is_selected": upload_method.is_selected,
            },
        }
    except HTTPException as exc:
        db.rollback()
        return error.InternalServerError(message=str(exc.detail))
    except Exception as e:
        db.rollback()
        logger.error(
            "Failed to update upload method selection uuid=%s: %s",
            upload_method_uuid,
            e,
        )
        return error.InternalServerError(message="Failed to update upload method selection")


# Retrieve documents
@router.get(
    "/{company_id}/documents",
    response_model=DocumentsResponse,
    response_model_exclude_none=True,
    tags=["rag"]
)
async def get_company_documents(
    request: Request,
    company_id: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    used_for: Optional[UploadUsedFor] = Query(
        None,
        description="Filter documents by usage context (`chat`, `voice_call`, `chat_and_voice_call`, or `market_planner`).",
    ),
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
    #api_key: str = Depends(api_key_security)

):

    try:
        
        # if not api_key or api_key != STATIC_API_KEY:
        #     raise HTTPException(status_code=401, detail="Unauthorized: Invalid API Key")
        
        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        base_query = (
            db.query(UploadMethod)
            .filter(UploadMethod.company_id == company_int_id)
            .order_by(UploadMethod.created_at.desc())
        )
        if used_for is not None:
            base_query = base_query.filter(
                UploadMethod.used_for.in_(
                    upload_method_used_for_values(used_for.value)
                )
            )
        
        documents, pagination = paginate_query(
            base_query,
            request,
            page,
            page_size
        )

        if used_for is not None:
            used_for_param = f"&used_for={used_for.value}"
            if pagination.get("next"):
                pagination["next"] += used_for_param
            if pagination.get("previous"):
                pagination["previous"] += used_for_param

        document_list = [
            DocumentItem(
                uuid=doc.uuid,
                filename=doc.filename,
                file_path=doc.file_path,
                is_selected=doc.is_selected if doc.is_selected is not None else False,
                used_for=doc.used_for,
                uploaded_at=doc.created_at
            )
            for doc in documents
        ]

        return {
            "success": True,
            "message": "Documents retrieved successfully",
            "data": {
                "documents": document_list,
                "pagination": pagination
            }
        }

    except ValueError as ve:
        return {
            "success": False,
            "error": str(ve)
        }

    except Exception as e:
        logger.error(f"Failed to retrieve documents: {str(e)}")
        return {
            "success": False,
            "error": "Failed to retrieve documents"
        }

@router.get(
    "/{company_id}/urls",
    response_model=UrlsResponse,
    response_model_exclude_none=True,
    tags=["rag"],
)
async def get_company_urls(
    request: Request,
    company_id: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    used_for: Optional[UploadUsedFor] = Query(
        None,
        description="Filter URLs by usage context (`chat`, `voice_call`, `chat_and_voice_call`, or `market_planner`).",
    ),
    db: Session = Depends(get_db),
    user: dict = Depends(auth.get_current_user),
):
    try:
        try:
            company_int_id = CompanyRepository(db).get_company_id_from_uuid(company_id)
        except HTTPException as exc:
            return error.InternalServerError(message=str(exc.detail))

        base_query = (
            db.query(UploadMethod)
            .filter(
                UploadMethod.company_id == company_int_id,
                UploadMethod.method_type == "url",
            )
            .order_by(UploadMethod.created_at.desc())
        )
        if used_for is not None:
            base_query = base_query.filter(
                UploadMethod.used_for.in_(
                    upload_method_used_for_values(used_for.value)
                )
            )

        url_records, pagination = paginate_query(
            base_query,
            request,
            page,
            page_size,
        )

        if used_for is not None:
            used_for_param = f"&used_for={used_for.value}"
            if pagination.get("next"):
                pagination["next"] += used_for_param
            if pagination.get("previous"):
                pagination["previous"] += used_for_param

        uploader_ids = {
            record.uploaded_by_id
            for record in url_records
            if record.uploaded_by_id is not None
        }
        # External user_id = users.uuid (same as user/roles/usage APIs)
        uploader_uuid_by_id = {}
        if uploader_ids:
            uploaders = db.query(User).filter(User.id.in_(uploader_ids)).all()
            uploader_uuid_by_id = {u.id: str(u.uuid) for u in uploaders}

        url_list = [
            UrlItem(
                uuid=record.uuid,
                url=record.url or record.file_path,
                filename=record.filename,
                is_selected=record.is_selected if record.is_selected is not None else False,
                used_for=record.used_for,
                user_id=uploader_uuid_by_id.get(record.uploaded_by_id),
                uploaded_at=record.created_at,
            )
            for record in url_records
        ]

        return {
            "success": True,
            "message": "URLs retrieved successfully",
            "data": {
                "urls": url_list,
                "pagination": pagination,
            },
        }

    except ValueError as ve:
        return {
            "success": False,
            "error": str(ve),
        }

    except Exception as e:
        logger.error(f"Failed to retrieve URLs: {str(e)}")
        return {
            "success": False,
            "error": "Failed to retrieve URLs",
        }


@router.delete(
    "/documents/{document_uuid}",
    tags=["rag"]
)
async def delete_document(
    document_uuid: str,
    db: Session = Depends(get_db)
):
    try:
        # Get document
        upload_method = RagRepository(db).get_upload_method_by_uuid(document_uuid)

        # Delete from S3
        if upload_method.method_type == "attachment" and upload_method.file_path:
            await delete_file_from_s3(upload_method.file_path)

        # Delete from DB (ExtractedData auto deletes via cascade)
        db.delete(upload_method)
        db.commit()

        return {
            "success": True,
            "message": "Document deleted successfully"
        }

    except HTTPException as exc:
        db.rollback()
        return error.InternalServerError(message=str(exc.detail))

    except Exception as e:
        db.rollback()
        logger.error(f"Failed to delete document: {str(e)}")
        return error.InternalServerError(message=f"Failed to delete document: {str(e)}")
           
